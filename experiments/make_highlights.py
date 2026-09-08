"""Generate the revised Highlights (Elsevier: 3-5 bullets, <=85 chars each) as .txt and .docx."""
from pathlib import Path

from docx import Document

ROOT = Path(r"E:/elsarticle-template-TMI_Revised")
R2 = ROOT / "revision2_materials"

HIGHLIGHTS = [
    "Endoscopic inter-frame motion is reformulated as pseudo-3D shape alignment.",
    "Gradient saliency is lifted to an interpretable pseudo-3D height field.",
    "A hybrid autoencoder learns deformation descriptors without manual annotation.",
    "Highlight-robust contour and most stable relative-rotation cue on weak-texture data.",
    "Calibrated stereo validation gives the best match precision and depth accuracy.",
]

for h in HIGHLIGHTS:
    n = len(h)
    status = "OK" if n <= 85 else "TOO LONG"
    print(f"[{n:2d} chars] {status}: {h}")
    assert n <= 85, f"bullet exceeds 85 chars: {h}"

out_dirs = [ROOT, ROOT / "revision2_materials"]
for out_dir in out_dirs:
    txt = out_dir / "Highlights_revised.txt"
    txt.write_text("Highlights\n\n" + "\n".join(f"- {h}" for h in HIGHLIGHTS) + "\n", encoding="utf-8")
    doc = Document()
    doc.add_heading("Highlights", level=1)
    for h in HIGHLIGHTS:
        doc.add_paragraph(h, style="List Bullet")
    doc.save(out_dir / "Highlights_revised.docx")
    print("written", txt)
